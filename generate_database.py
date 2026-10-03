import openpyxl
import json

wb = openpyxl.load_workbook("recipe.xlsx", data_only=True)

def unmerge_and_fill(ws):
    for mrange in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = mrange.min_col, mrange.min_row, mrange.max_col, mrange.max_row
        top_left_val = ws.cell(min_row, min_col).value
        ws.unmerge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(r, c).value = top_left_val

for sname in ["오리지널 티", "과일티", "밀크티", "라떼", "치즈 밀크폼"]:
    unmerge_and_fill(wb[sname])

# Shorten Form
ws_s = wb["쇼튼 폼"]
liquid_map = {}
for r in [3, 5, 7, 9, 11, 13]:
    for c in range(2, ws_s.max_column+1):
        t, b = ws_s.cell(r, c).value, ws_s.cell(r+1, c).value
        if t is not None and b is not None:
            liquid_map[float(t)] = float(b)

powder_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(17, c).value, ws_s.cell(18, c).value
    if t is not None and b is not None:
        powder_map[float(t)] = float(b)

ice_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(21, c).value, ws_s.cell(22, c).value
    if t is not None and b is not None:
        ice_map[float(t)] = float(b)

# Helper to parse syrup string e.g. "40 25 15 10" -> {100: 40, 50: 25, 30: 15, 10: 10, 0: 0}
def parse_syrup(val):
    if not val:
        return {100: 0, 50: 0, 30: 0, 10: 0, 0: 0}
    if isinstance(val, (int, float)):
        return {100: float(val), 50: float(val), 30: float(val), 10: float(val), 0: 0}
    parts = [float(x) for x in str(val).strip().split()]
    if len(parts) == 4:
        return {100: parts[0], 50: parts[1], 30: parts[2], 10: parts[3], 0: 0}
    elif len(parts) == 3:
        return {100: parts[0], 50: parts[1], 30: parts[2], 10: 0, 0: 0}
    elif len(parts) == 1:
        return {100: parts[0], 50: parts[0], 30: parts[0], 10: parts[0], 0: 0}
    return {100: 0, 50: 0, 30: 0, 10: 0, 0: 0}

database = {
    "shorten": {
        "liquid": liquid_map,
        "powder": powder_map,
        "ice": ice_map
    },
    "categories": {}
}

# 1. 오리지널 티
ws = wb["오리지널 티"]
ice_cols = {
    "얼음 보통": 4,
    "얼음 적게": 5,
    "매우적게": 6,
    "얼음 없이": 7,
    "뜨겁게": 8
}

orig_items = [
    {
        "names": ["루이보스", "루이보스 티", "루이보스 우롱티"],
        "display": "루이보스",
        "tea_name": "루이보스 티",
        "rows": {"M": 4, "L": 7, "count": 3},
        "has_water": False
    },
    {
        "names": ["블랙", "블랙 우롱티", "블랙 티"],
        "display": "블랙 우롱티",
        "tea_name": "블랙 우롱티",
        "rows": {"M": 10, "L": 14, "count": 4},
        "has_water": True
    },
    {
        "names": ["그린", "그린 우롱티"],
        "display": "그린 우롱티",
        "tea_name": "그린 우롱티",
        "rows": {"M": 18, "L": 21, "count": 3},
        "has_water": False
    },
    {
        "names": ["라이트", "라이트 우롱티"],
        "display": "라이트 우롱티",
        "tea_name": "라이트 우롱티",
        "rows": {"M": 18, "L": 21, "count": 3},
        "has_water": False
    },
    {
        "names": ["다크", "다크 우롱티", "다크 로스티드 우롱티"],
        "display": "다크 우롱티",
        "tea_name": "다크 우롱티",
        "rows": {"M": 18, "L": 21, "count": 3},
        "has_water": False
    },
    {
        "names": ["스프링", "스프링 우롱티"],
        "display": "스프링 우롱티",
        "tea_name": "스프링 우롱티",
        "rows": {"M": 24, "L": 27, "count": 3},
        "has_water": False
    }
]

cat_orig = {
    "name": "오리지널 티",
    "order_rule": "얼음 → 시럽 → 티 (뜨겁게: 시럽 → 티 / 블랙: 온수 추가)",
    "ice_options": ["얼음 보통", "얼음 적게", "매우적게", "얼음 없이", "뜨겁게"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

for item in orig_items:
    item_data = {"display": item["display"], "recipes": {}}
    for size in ["M", "L"]:
        item_data["recipes"][size] = {}
        start_r = item["rows"][size]
        for ice_opt, col in ice_cols.items():
            # row 0: 얼음
            ice_val = ws.cell(start_r, col).value
            if ice_val == "x" or ice_val is None:
                ice_amt = 0.0
            else:
                ice_amt = float(ice_val)
            
            # row 1: 시럽
            syrup_map = parse_syrup(ws.cell(start_r + 1, col).value)
            
            # row 2: 티
            tea_amt = float(ws.cell(start_r + 2, col).value or 0)
            
            rec = {
                "얼음": ice_amt,
                "시럽": syrup_map,
                item["tea_name"]: tea_amt
            }
            if item["has_water"]:
                water_amt = float(ws.cell(start_r + 3, col).value or 0)
                rec["온수"] = water_amt
            
            item_data["recipes"][size][ice_opt] = rec
            
    cat_orig["items"][item["display"]] = item_data

database["categories"]["오리지널 티"] = cat_orig
print("오리지널 티 parsed.")
