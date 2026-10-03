import openpyxl
import json

wb = openpyxl.load_workbook("recipe.xlsx", data_only=True)

# 1. Parse Shorten Form
ws_s = wb["쇼튼 폼"]
liquid_map = {}
for r in [3, 5, 7, 9, 11, 13]:
    for c in range(2, ws_s.max_column+1):
        t, b = ws_s.cell(r, c).value, ws_s.cell(r+1, c).value
        if t is not None and b is not None:
            liquid_map[float(t)] = float(b)

powder_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(17, c).value, ws_s.cell(18, c).value
    if t is not None and b is not None:
        powder_map[float(t)] = float(b)

ice_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(21, c).value, ws_s.cell(22, c).value
    if t is not None and b is not None:
        ice_map[float(t)] = float(b)

shorten_tables = {
    "liquid": liquid_map,
    "powder": powder_map,
    "ice": ice_map
}

with open("shorten_data.json", "w", encoding="utf-8") as f:
    json.dump(shorten_tables, f, ensure_ascii=False, indent=2)

print("Shorten tables exported successfully.")
