import openpyxl
import json

wb = openpyxl.load_workbook("recipe.xlsx", data_only=True)

def unmerge_and_fill(ws):
    for mrange in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = mrange.min_col, mrange.min_row, mrange.max_col, mrange.max_row
        top_left_val = ws.cell(min_row, min_col).value
        ws.unmerge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(r, c).value = top_left_val

for sname in ["오리지널 티", "과일티", "밀크티", "라떼", "치즈 밀크폼"]:
    unmerge_and_fill(wb[sname])

# Shorten Form
ws_s = wb["쇼튼 폼"]
liquid_map = {}
for r in [3, 5, 7, 9, 11, 13]:
    for c in range(2, ws_s.max_column+1):
        t, b = ws_s.cell(r, c).value, ws_s.cell(r+1, c).value
        if t is not None and b is not None:
            liquid_map[float(t)] = float(b)

powder_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(17, c).value, ws_s.cell(18, c).value
    if t is not None and b is not None:
        powder_map[float(t)] = float(b)

ice_map = {}
for c in range(2, ws_s.max_column+1):
    t, b = ws_s.cell(21, c).value, ws_s.cell(22, c).value
    if t is not None and b is not None:
        ice_map[float(t)] = float(b)

def parse_syrup(val):
    if not val:
        return {"100%": 0.0, "50%": 0.0, "30%": 0.0, "10%": 0.0, "0%": 0.0}
    if isinstance(val, (int, float)):
        return {"100%": float(val), "50%": float(val), "30%": float(val), "10%": float(val), "0%": 0.0}
    parts = [float(x) for x in str(val).strip().split()]
    if len(parts) == 4:
        return {"100%": parts[0], "50%": parts[1], "30%": parts[2], "10%": parts[3], "0%": 0.0}
    elif len(parts) == 3:
        return {"100%": parts[0], "50%": parts[1], "30%": parts[2], "10%": 0.0, "0%": 0.0}
    return {"100%": 0.0, "50%": 0.0, "30%": 0.0, "10%": 0.0, "0%": 0.0}

database = {
    "shorten": {
        "liquid": liquid_map,
        "powder": powder_map,
        "ice": ice_map
    },
    "categories": {}
}

# 1. 오리지널 티
ws_orig = wb["오리지널 티"]
ice_cols_5 = {"얼음 보통": 4, "얼음 적게": 5, "매우적게": 6, "얼음 없이": 7, "뜨겁게": 8}
orig_items = [
    {"display": "루이보스", "sheet_tea": "루이보스", "tea_name": "루이보스 티", "rows": {"M": 4, "L": 7}, "has_water": False},
    {"display": "블랙 우롱티", "sheet_tea": "블랙", "tea_name": "블랙 티", "rows": {"M": 10, "L": 14}, "has_water": True},
    {"display": "그린 우롱티", "sheet_tea": "그린", "tea_name": "그린 티", "rows": {"M": 18, "L": 21}, "has_water": False},
    {"display": "라이트 우롱티", "sheet_tea": "라이트", "tea_name": "라이트 티", "rows": {"M": 18, "L": 21}, "has_water": False},
    {"display": "다크 우롱티", "sheet_tea": "다크", "tea_name": "다크 티", "rows": {"M": 18, "L": 21}, "has_water": False},
    {"display": "스프링 우롱티", "sheet_tea": "스프링", "tea_name": "스프링 티", "rows": {"M": 24, "L": 27}, "has_water": False}
]

cat_orig = {
    "name": "오리지널 티",
    "order_rule": "얼음 → 시럽 → 티 (뜨겁게: 시럽 → 티 / 블랙: 온수 추가)",
    "ice_options": ["얼음 보통", "얼음 적게", "매우적게", "얼음 없이", "뜨겁게"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

for item in orig_items:
    item_data = {"display": item["display"], "recipes": {}}
    for size in ["M", "L"]:
        item_data["recipes"][size] = {}
        start_r = item["rows"][size]
        for ice_opt, col in ice_cols_5.items():
            ice_val = ws_orig.cell(start_r, col).value
            ice_amt = 0.0 if (ice_val == "x" or ice_val is None) else float(ice_val)
            syrup_map = parse_syrup(ws_orig.cell(start_r + 1, col).value)
            tea_amt = float(ws_orig.cell(start_r + 2, col).value or 0)
            
            # Step sequence in order
            steps = []
            if ice_opt != "뜨겁게" and ice_amt > 0:
                steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
            steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
            steps.append({"name": item["tea_name"], "type": "liquid", "unit": "ml", "amount": tea_amt})
            if item["has_water"]:
                water_amt = float(ws_orig.cell(start_r + 3, col).value or 0)
                steps.append({"name": "온수", "type": "liquid", "unit": "ml", "amount": water_amt})
            
            item_data["recipes"][size][ice_opt] = {
                "steps": steps,
                "ice": ice_amt,
                "syrup": syrup_map,
                "tea": tea_amt,
                "water": float(ws_orig.cell(start_r + 3, col).value or 0) if item["has_water"] else 0
            }
    cat_orig["items"][item["display"]] = item_data
database["categories"]["오리지널 티"] = cat_orig

# 2. 과일티
ws_fruit = wb["과일티"]
fruit_ice_cols = {"얼음 보통": 4, "얼음 적게": 5, "매우적게": 6, "얼음 없이": 7}
cat_fruit = {
    "name": "과일티",
    "order_rule": "얼음 → 시럽 → 주스 → 티",
    "ice_options": ["얼음 보통", "얼음 적게", "매우적게", "얼음 없이"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

# 자몽
jamong_recipes = {}
for size, sr in [("M", 4), ("L", 9)]:
    jamong_recipes[size] = {}
    for ice_opt, col in fruit_ice_cols.items():
        ice_val = ws_fruit.cell(sr, col).value
        ice_amt = 0.0 if (ice_val is None or ice_val == "x" or ice_opt == "얼음 없이") else float(ice_val)
        syrup_map = parse_syrup(ws_fruit.cell(sr+1, col).value)
        jamong_juice = float(ws_fruit.cell(sr+2, col).value or 0)
        pomelo_juice = float(ws_fruit.cell(sr+3, col).value or 0)
        dark_tea = float(ws_fruit.cell(sr+4, col).value or 0)
        
        steps = []
        if ice_amt > 0:
            steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
        steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
        steps.append({"name": "자몽 주스", "type": "liquid", "unit": "ml", "amount": jamong_juice})
        steps.append({"name": "포멜로 주스", "type": "liquid", "unit": "ml", "amount": pomelo_juice})
        steps.append({"name": "다크 티", "type": "liquid", "unit": "ml", "amount": dark_tea})
        
        jamong_recipes[size][ice_opt] = {
            "steps": steps,
            "ice": ice_amt,
            "syrup": syrup_map,
            "자몽 주스": jamong_juice,
            "포멜로 주스": pomelo_juice,
            "다크 티": dark_tea
        }
cat_fruit["items"]["자몽 우롱티"] = {"display": "자몽 우롱티", "recipes": jamong_recipes}

# 오렌지
orange_recipes = {}
for size, sr in [("M", 15), ("L", 21)]:
    orange_recipes[size] = {}
    for ice_opt, col in fruit_ice_cols.items():
        ice_val = ws_fruit.cell(sr, col).value
        ice_amt = 0.0 if (ice_val is None or ice_val == "x" or ice_opt == "얼음 없이") else float(ice_val)
        syrup_map = parse_syrup(ws_fruit.cell(sr+1, col).value)
        orange_juice = float(ws_fruit.cell(sr+2, col).value or 0)
        spring_tea = float(ws_fruit.cell(sr+3, col).value or 0)
        dark_tea = float(ws_fruit.cell(sr+5, col).value or 0)
        
        steps = []
        if ice_amt > 0:
            steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
        steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
        steps.append({"name": "오렌지 주스", "type": "liquid", "unit": "ml", "amount": orange_juice})
        steps.append({"name": "스프링 티", "type": "liquid", "unit": "ml", "amount": spring_tea, "zero_sugar_extra": 30.0})
        steps.append({"name": "다크 티", "type": "liquid", "unit": "ml", "amount": dark_tea})
        
        orange_recipes[size][ice_opt] = {
            "steps": steps,
            "ice": ice_amt,
            "syrup": syrup_map,
            "오렌지 주스": orange_juice,
            "스프링 티": spring_tea,
            "다크 티": dark_tea,
            "special_rule": "당도 0% 선택 시: 스프링 우롱티 +30ml"
        }
cat_fruit["items"]["오렌지 우롱티"] = {"display": "오렌지 우롱티", "recipes": orange_recipes}

# 레몬
lemon_recipes = {}
for size, sr in [("M", 28), ("L", 34)]:
    lemon_recipes[size] = {}
    for ice_opt, col in fruit_ice_cols.items():
        ice_val = ws_fruit.cell(sr, col).value
        ice_amt = 0.0 if (ice_val is None or ice_val == "x" or ice_opt == "얼음 없이") else float(ice_val)
        syrup_map = parse_syrup(ws_fruit.cell(sr+1, col).value)
        lemon_juice = float(ws_fruit.cell(sr+2, col).value or 0)
        spring_tea = float(ws_fruit.cell(sr+3, col).value or 0)
        dark_tea = float(ws_fruit.cell(sr+4, col).value or 0)
        
        steps = []
        if ice_amt > 0:
            steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
        steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
        steps.append({"name": "레몬 주스", "type": "liquid", "unit": "ml", "amount": lemon_juice})
        steps.append({"name": "스프링 티", "type": "liquid", "unit": "ml", "amount": spring_tea})
        steps.append({"name": "다크 티", "type": "liquid", "unit": "ml", "amount": dark_tea, "zero_sugar_extra": 30.0})
        
        lemon_recipes[size][ice_opt] = {
            "steps": steps,
            "ice": ice_amt,
            "syrup": syrup_map,
            "레몬 주스": lemon_juice,
            "스프링 티": spring_tea,
            "다크 티": dark_tea,
            "special_rule": "당도 0% 선택 시: 다크 우롱티 +30ml"
        }
cat_fruit["items"]["레몬 우롱티"] = {"display": "레몬 우롱티", "recipes": lemon_recipes}

database["categories"]["과일티"] = cat_fruit

# 3. 밀크티
ws_milk = wb["밀크티"]
cat_milk = {
    "name": "밀크티",
    "order_rule": "ICE: 크리머 → 시럽 → 티 → 얼음 / HOT: 크리머 → 시럽 → 티",
    "ice_options": ["얼음 보통", "얼음 적게", "매우적게", "얼음 없이", "뜨겁게"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

milk_teas = [
    {"display": "블랙 우롱 밀크티", "tea_name": "블랙 티"},
    {"display": "그린 우롱 밀크티", "tea_name": "그린 티"},
    {"display": "라이트 우롱 밀크티", "tea_name": "라이트 티"},
    {"display": "다크 로스티드 우롱 밀크티", "tea_name": "다크 티"}
]

for mtea in milk_teas:
    mtea_recipes = {}
    for size, sr in [("M", 4), ("L", 8)]:
        mtea_recipes[size] = {}
        for ice_opt, col in ice_cols_5.items():
            ice_val = ws_milk.cell(sr, col).value
            ice_amt = 0.0 if (ice_val == "x" or ice_val is None) else float(ice_val)
            syrup_map = parse_syrup(ws_milk.cell(sr+1, col).value)
            creamer = float(ws_milk.cell(sr+2, col).value or 0)
            tea_amt = float(ws_milk.cell(sr+3, col).value or 0)
            
            steps = []
            steps.append({"name": "크리머", "type": "powder", "unit": "스쿱", "amount": creamer})
            steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
            steps.append({"name": mtea["tea_name"], "type": "liquid", "unit": "ml", "amount": tea_amt})
            if ice_opt != "뜨겁게" and ice_amt > 0:
                steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
            
            mtea_recipes[size][ice_opt] = {
                "steps": steps,
                "ice": ice_amt,
                "syrup": syrup_map,
                "creamer": creamer,
                "tea": tea_amt
            }
    cat_milk["items"][mtea["display"]] = {"display": mtea["display"], "recipes": mtea_recipes}

# 호지차 밀크티
hoji_milk_recipes = {}
for size, sr in [("M", 12), ("L", 17)]:
    hoji_milk_recipes[size] = {}
    for ice_opt, col in ice_cols_5.items():
        ice_val = ws_milk.cell(sr, col).value
        ice_amt = 0.0 if (ice_val == "x" or ice_val is None) else float(ice_val)
        syrup_map = parse_syrup(ws_milk.cell(sr+1, col).value)
        hoji_powder = float(ws_milk.cell(sr+2, col).value or 0)
        creamer = float(ws_milk.cell(sr+3, col).value or 0)
        hot_water = float(ws_milk.cell(sr+4, col).value or 0)
        
        steps = []
        steps.append({"name": "크리머", "type": "powder", "unit": "스쿱", "amount": creamer})
        steps.append({"name": "호지차 파우더", "type": "powder", "unit": "스쿱", "amount": hoji_powder})
        steps.append({"name": "온수", "type": "liquid", "unit": "ml", "amount": hot_water, "note": "파우더 녹이기"})
        steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
        if ice_opt != "뜨겁게" and ice_amt > 0:
            steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
            
        hoji_milk_recipes[size][ice_opt] = {
            "steps": steps,
            "ice": ice_amt,
            "syrup": syrup_map,
            "호지차 파우더": hoji_powder,
            "크리머": creamer,
            "온수": hot_water
        }
cat_milk["items"]["호지차 우롱 밀크티"] = {"display": "호지차 우롱 밀크티", "recipes": hoji_milk_recipes}

database["categories"]["밀크티"] = cat_milk

# 4. 라떼
ws_latte = wb["라떼"]
cat_latte = {
    "name": "라떼",
    "order_rule": "ICE: 얼음 → 시럽 → 티 → 시럽 녹을 때까지 젓기 → 우유 → 얼음/상온수 추가 / HOT: 시럽 → 티 → 우유",
    "ice_options": ["얼음 보통", "얼음 적게", "매우적게", "얼음 없이", "뜨겁게"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

latte_teas = [
    {"display": "블랙 우롱 라떼", "tea_name": "블랙 티"},
    {"display": "그린 우롱 라떼", "tea_name": "그린 티"},
    {"display": "라이트 우롱 라떼", "tea_name": "라이트 티"},
    {"display": "다크 우롱 라떼", "tea_name": "다크 티"}
]

for ltea in latte_teas:
    ltea_recipes = {}
    for size, sr in [("M", 4), ("L", 8)]:
        ltea_recipes[size] = {}
        for ice_opt, col in ice_cols_5.items():
            ice_val = ws_latte.cell(sr, col).value
            ice_amt = 0.0 if (ice_val == "x" or ice_val is None) else float(ice_val)
            syrup_map = parse_syrup(ws_latte.cell(sr+1, col).value)
            tea_amt = float(ws_latte.cell(sr+2, col).value or 0)
            milk_amt = float(ws_latte.cell(sr+3, col).value or 0)
            
            steps = []
            if ice_opt != "뜨겁게" and ice_amt > 0:
                steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
            steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
            steps.append({"name": ltea["tea_name"], "type": "liquid", "unit": "ml", "amount": tea_amt, "action": "시럽 녹을 때까지 젓기"})
            steps.append({"name": "우유", "type": "liquid", "unit": "ml", "amount": milk_amt})
            if ice_opt != "뜨겁게":
                steps.append({"name": "얼음/상온수 채우기", "type": "action", "unit": "", "amount": 0, "note": "선까지 채우기"})
            
            ltea_recipes[size][ice_opt] = {
                "steps": steps,
                "ice": ice_amt,
                "syrup": syrup_map,
                "tea": tea_amt,
                "milk": milk_amt
            }
    cat_latte["items"][ltea["display"]] = {"display": ltea["display"], "recipes": ltea_recipes}

# 호지차 라떼
hoji_latte_recipes = {}
for size, sr in [("M", 13), ("L", 18)]:
    hoji_latte_recipes[size] = {}
    for ice_opt, col in ice_cols_5.items():
        ice_val = ws_latte.cell(sr, col).value
        ice_amt = 0.0 if (ice_val == "x" or ice_val is None) else float(ice_val)
        syrup_map = parse_syrup(ws_latte.cell(sr+1, col).value)
        hoji_powder = float(ws_latte.cell(sr+2, col).value or 0)
        hot_water = float(ws_latte.cell(sr+3, col).value or 0)
        milk_amt = float(ws_latte.cell(sr+4, col).value or 0)
        
        steps = []
        if ice_opt != "뜨겁게" and ice_amt > 0:
            steps.append({"name": "얼음", "type": "ice", "unit": "스쿱", "amount": ice_amt})
        steps.append({"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map})
        steps.append({"name": "호지차 파우더", "type": "powder", "unit": "스쿱", "amount": hoji_powder})
        steps.append({"name": "뜨거운 물", "type": "liquid", "unit": "ml", "amount": hot_water, "note": "파우더 녹이기"})
        steps.append({"name": "우유", "type": "liquid", "unit": "ml", "amount": milk_amt})
        if ice_opt != "뜨겁게":
            steps.append({"name": "얼음/상온수 채우기", "type": "action", "unit": "", "amount": 0, "note": "선까지 채우기"})
            
        hoji_latte_recipes[size][ice_opt] = {
            "steps": steps,
            "ice": ice_amt,
            "syrup": syrup_map,
            "호지차 파우더": hoji_powder,
            "뜨거운 물": hot_water,
            "우유": milk_amt
        }
cat_latte["items"]["호지차 라떼"] = {"display": "호지차 라떼", "recipes": hoji_latte_recipes}

database["categories"]["라떼"] = cat_latte

# 5. 치즈 밀크폼
ws_cheese = wb["치즈 밀크폼"]
cat_cheese = {
    "name": "치즈 밀크폼",
    "order_rule": "티 → 얼음 → 시럽 (+ 치즈 밀크폼 얹기)",
    "ice_options": ["얼음 보통"],
    "sugar_options": ["100%", "50%", "30%", "10%", "0%"],
    "items": {}
}

# 블랙
cheese_black = {}
for size, sr in [("M", 4), ("L", 8)]:
    syrup_map = parse_syrup(ws_cheese.cell(sr+1, 4).value)
    steps = [
        {"name": "블랙 티", "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+2, 4).value or 0)},
        {"name": "온수", "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+3, 4).value or 0)},
        {"name": "얼음", "type": "ice", "unit": "스쿱", "amount": float(ws_cheese.cell(sr, 4).value or 0)},
        {"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map},
        {"name": "치즈 밀크폼", "type": "topping_cream", "unit": "ml", "amount": 0, "note": "상단 가득 붓기"}
    ]
    cheese_black[size] = {"얼음 보통": {"steps": steps}}
cat_cheese["items"]["치즈 밀크폼 블랙 우롱티"] = {"display": "치즈 밀크폼 블랙 우롱티", "recipes": cheese_black}

# 그린, 라이트, 다크
for tea_disp, tea_sub in [("치즈 밀크폼 그린 우롱티", "그린 티"), ("치즈 밀크폼 라이트 우롱티", "라이트 티"), ("치즈 밀크폼 다크 우롱티", "다크 티")]:
    c_rec = {}
    for size, sr in [("M", 12), ("L", 16)]:
        syrup_map = parse_syrup(ws_cheese.cell(sr+1, 4).value)
        steps = [
            {"name": tea_sub, "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+2, 4).value or 0)},
            {"name": "얼음", "type": "ice", "unit": "스쿱", "amount": float(ws_cheese.cell(sr, 4).value or 0)},
            {"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map},
            {"name": "치즈 밀크폼", "type": "topping_cream", "unit": "", "amount": 0, "note": "상단 가득 붓기"}
        ]
        c_rec[size] = {"얼음 보통": {"steps": steps}}
    cat_cheese["items"][tea_disp] = {"display": tea_disp, "recipes": c_rec}

# 스프링
c_spring = {}
for size, sr in [("M", 12), ("L", 16)]:
    syrup_map = parse_syrup(ws_cheese.cell(sr+1, 4).value)
    steps = [
        {"name": "스프링 티", "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+3, 4).value or 0)},
        {"name": "얼음", "type": "ice", "unit": "스쿱", "amount": float(ws_cheese.cell(sr, 4).value or 0)},
        {"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map},
        {"name": "치즈 밀크폼", "type": "topping_cream", "unit": "", "amount": 0, "note": "상단 가득 붓기"}
    ]
    c_spring[size] = {"얼음 보통": {"steps": steps}}
cat_cheese["items"]["치즈 밀크폼 스프링 우롱티"] = {"display": "치즈 밀크폼 스프링 우롱티", "recipes": c_spring}

# 초코
c_choco = {}
for size, sr in [("M", 20), ("L", 24)]:
    syrup_map = parse_syrup(ws_cheese.cell(sr+1, 4).value)
    steps = [
        {"name": "초코 파우더", "type": "powder", "unit": "스쿱", "amount": float(ws_cheese.cell(sr+2, 4).value or 0)},
        {"name": "온수", "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+3, 4).value or 0), "note": "파우더 녹이기"},
        {"name": "얼음", "type": "ice", "unit": "스쿱", "amount": float(ws_cheese.cell(sr, 4).value or 0)},
        {"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map},
        {"name": "치즈 밀크폼", "type": "topping_cream", "unit": "", "amount": 0, "note": "상단 가득 붓기"}
    ]
    c_choco[size] = {"얼음 보통": {"steps": steps}}
cat_cheese["items"]["치즈 밀크폼 초코"] = {"display": "치즈 밀크폼 초코", "recipes": c_choco}

# 호지차
c_hoji = {}
for size, sr in [("M", 28), ("L", 33)]:
    syrup_map = parse_syrup(ws_cheese.cell(sr+1, 4).value)
    steps = [
        {"name": "호지차 파우더", "type": "powder", "unit": "스쿱", "amount": float(ws_cheese.cell(sr+2, 4).value or 0)},
        {"name": "크리머", "type": "powder", "unit": "스쿱", "amount": float(ws_cheese.cell(sr+3, 4).value or 0)},
        {"name": "온수", "type": "liquid", "unit": "ml", "amount": float(ws_cheese.cell(sr+4, 4).value or 0), "note": "파우더 녹이기"},
        {"name": "얼음", "type": "ice", "unit": "스쿱", "amount": float(ws_cheese.cell(sr, 4).value or 0)},
        {"name": "시럽", "type": "liquid", "unit": "cc", "amount_by_sugar": syrup_map},
        {"name": "치즈 밀크폼", "type": "topping_cream", "unit": "", "amount": 0, "note": "상단 가득 붓기"}
    ]
    c_hoji[size] = {"얼음 보통": {"steps": steps}}
cat_cheese["items"]["치즈 밀크폼 호지차"] = {"display": "치즈 밀크폼 호지차", "recipes": c_hoji}

database["categories"]["치즈 밀크폼"] = cat_cheese

with open("recipe_database.json", "w", encoding="utf-8") as f:
    json.dump(database, f, ensure_ascii=False, indent=2)

with open("recipe_database.js", "w", encoding="utf-8") as f:
    f.write("window.RECIPE_DATABASE = " + json.dumps(database, ensure_ascii=False, indent=2) + ";\n")

print("Full database generated successfully!")
